"""Build the JH&P demo takeaway deck from the TMF takeaway deck's template.

The TMF deck supplies the masters, layouts and the five slide designs we reuse:
cover, requirements table, chapter divider, screenshot (with optional callout)
and thank-you. Each new slide is a deep copy of one of those prototypes with
its text and pictures swapped; the original TMF slides are then removed.

Run:  python3 build_deck.py spec.json out.pptx
"""
import copy, json, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEMPLATE = '/Users/lucia.ribot/Downloads/TMF _ Unified Client Portal _ DEMO Takeaway Deck.pptx'
PROTO = {'cover': 0, 'table': 1, 'divider': 2, 'shot': 6, 'thanks': 59}

spec = json.load(open(sys.argv[1]))
prs = Presentation(TEMPLATE)
orig = list(prs.slides)
n_orig = len(orig)

def clone(src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('/image') or rel.reltype.endswith('/hyperlink'):
            if rel.is_external:
                rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in (qn('r:embed'), qn('r:link'), qn('r:id')):
                v = node.get(attr)
                if v in rid_map:
                    node.set(attr, rid_map[v])
        new.shapes._spTree.append(el)
    # keep the background (cover uses a picture placeholder, others inherit)
    return new

def set_text(shape, text, keep_first_run=True):
    """Replace a shape's text, keeping the first paragraph/run formatting."""
    tf = shape.text_frame
    paras = text.split('\n') if isinstance(text, str) else text
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else None
    rpr = copy.deepcopy(r0._r.find(qn('a:rPr'))) if r0 is not None and r0._r.find(qn('a:rPr')) is not None else None
    ppr = copy.deepcopy(p0._p.find(qn('a:pPr'))) if p0._p.find(qn('a:pPr')) is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        r = p.add_run(); r.text = line
        if rpr is not None:
            old = r._r.find(qn('a:rPr'))
            if old is not None: r._r.remove(old)
            r._r.insert(0, copy.deepcopy(rpr))

def by_name(slide, name):
    for s in slide.shapes:
        if s.name == name: return s
    raise KeyError(name)

def swap_picture(slide, shape, image_path):
    """Point an existing picture (or filled picture placeholder) at a new image."""
    blip = shape._element.find('.//' + qn('a:blip'))
    img_part, rid = slide.part.get_or_add_image_part(image_path)
    blip.set(qn('r:embed'), rid)

def logo(slide, path):
    for s in slide.shapes:
        if s.shape_type == 13 and s.left == 8340400 and s.top == 213650:
            swap_picture(slide, s, path)

made = []
for item in spec['slides']:
    kind = item['kind']
    src = orig[PROTO[kind]]
    s = clone(src)
    made.append(s)
    if kind == 'cover':
        set_text(by_name(s, 'Google Shape;417;p67'), item['title'])
        set_text(by_name(s, 'Google Shape;418;p67'), item['people'])
        set_text(by_name(s, 'Google Shape;419;p67'), item['date'])
        logo_shape = by_name(s, 'Google Shape;420;p67'); swap_picture(s, logo_shape, spec['logo'])
    elif kind == 'table':
        set_text(by_name(s, 'Google Shape;426;p68'), item['title'])
        pic = by_name(s, 'Google Shape;427;p68'); pic._element.getparent().remove(pic._element)
        rows, cols = len(item['rows']) + 1, len(item['header'])
        x, y, w = Emu(402838), Emu(1020000), Emu(8338200)
        tbl = s.shapes.add_table(rows, cols, x, y, w, Emu(300000 * rows)).table
        widths = item.get('widths', [1 / cols] * cols)
        for i, f in enumerate(widths): tbl.columns[i].width = Emu(int(8338200 * f))
        def cell(r, c, text, bold=False, color=None, fill=None, size=8):
            ce = tbl.cell(r, c); ce.text = ''
            p = ce.text_frame.paragraphs[0]; run = p.add_run(); run.text = text
            run.font.size = Pt(size); run.font.bold = bold
            if color: run.font.color.rgb = RGBColor.from_string(color)
            ce.margin_left = ce.margin_right = Emu(60000); ce.margin_top = ce.margin_bottom = Emu(30000)
            if fill: ce.fill.solid(); ce.fill.fore_color.rgb = RGBColor.from_string(fill)
        for c, h in enumerate(item['header']): cell(0, c, h, True, 'FFFFFF', '24294F', 9)
        for r, row in enumerate(item['rows'], start=1):
            for c, v in enumerate(row): cell(r, c, v, c == 0, '1F2437', 'FFFFFF' if r % 2 else 'F2ECE7')
        logo(s, spec['logo'])
    elif kind == 'divider':
        body = by_name(s, 'Google Shape;433;p69')
        tf = body.text_frame
        # paragraphs: '', number, title+description runs — rewrite run text in place
        texts = [item['number'], item['title'], item['description']]
        runs = [r for p in tf.paragraphs for r in p.runs if r.text.strip()]
        for r, t in zip(runs, texts): r.text = t
        for r in runs[len(texts):]: r.text = ''
    elif kind == 'shot':
        tf = by_name(s, 'Google Shape;465;p73').text_frame
        p0, p1 = tf.paragraphs[0], tf.paragraphs[1]
        p0.runs[0].text = item['title']
        for r in p0.runs[1:]: r.text = ''
        if item.get('subtitle'):
            p1.runs[0].text = item['subtitle']
            for r in p1.runs[1:]: r.text = ''
        else:
            p1._p.getparent().remove(p1._p)
        pic = by_name(s, 'Google Shape;463;p73')
        swap_picture(s, pic, item['image'])
        for sr in pic._element.iter(qn('a:srcRect')): sr.getparent().remove(sr)
        box = by_name(s, 'Google Shape;466;p73'); txt = by_name(s, 'Google Shape;467;p73')
        if item.get('callout'):
            set_text(txt, item['callout'])
            for r in txt.text_frame.paragraphs[0].runs: r.font.size = Pt(10); r.font.name = 'Arial'; r.font.color.rgb = RGBColor(0x1F, 0x24, 0x37)
            W = Emu(1650000)
            side = item.get('side', 'left')
            box.width = W; txt.width = Emu(W - 140000)
            box.height = Emu(1500000); txt.height = Emu(1300000)
            if side == 'right':
                pic.left = Emu(1520000)
                box.left = Emu(pic.left + pic.width + 120000)
            else:
                pic.left = Emu(106250 + W + 120000)
                box.left = Emu(106250)
            txt.left = Emu(box.left + 70000)
            box.top = Emu(item.get('y', 2200000)); txt.top = Emu(box.top + 100000)
        else:
            box._element.getparent().remove(box._element); txt._element.getparent().remove(txt._element)
        logo(s, spec['logo'])
    elif kind == 'thanks':
        pass

# drop the original TMF slides
sldIdLst = prs.slides._sldIdLst
for sl in orig:
    for sldId in list(sldIdLst):
        if prs.part.related_part(sldId.rId) is sl.part:
            prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
prs.save(sys.argv[2])
print('saved', sys.argv[2], len(prs.slides), 'slides')
